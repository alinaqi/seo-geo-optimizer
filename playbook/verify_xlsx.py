"""Post-build checks: sheets, source ID integrity, probability ranges, formula values after recalc."""
import re, sys
from openpyxl import load_workbook

EXPECTED = ["README", "Executive Summary", "Strategy Playbook", "Agent Execution", "Validation Checklist",
            "Penalty Safeguards", "Anti-patterns", "Agent Test Scenarios", "SEO Tools", "AEO Tools",
            "Measurement KPIs", "Prompt Set Design", "Timeline Aug-Oct 2026", "Sources"]


def source_ids(wb):
    ws = wb["Sources"]
    return {ws.cell(row=r, column=1).value for r in range(5, ws.max_row + 1) if ws.cell(row=r, column=1).value}


def referenced_ids(wb):
    refs = set()
    for ws in wb.worksheets:
        if ws.title == "Sources":
            continue
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str):
                    refs.update(re.findall(r"\bS\d{2}\b", v))
    return refs


def check_strategies(ws):
    probs = [ws.cell(row=r, column=8).value for r in range(5, ws.max_row + 1) if ws.cell(row=r, column=1).value and str(ws.cell(row=r, column=1).value).startswith("ST")]
    assert all(isinstance(p, (int, float)) and 0 < p <= 1 for p in probs), probs
    return len(probs)


def check_values(path):
    wb = load_workbook(path, data_only=True)
    ws = wb["Strategy Playbook"]
    vals = [ws.cell(row=r, column=12).value for r in range(5, 5 + 24)]
    assert all(isinstance(v, (int, float)) for v in vals), vals
    es = wb["Executive Summary"]
    stats = [es.cell(row=r, column=2).value for r in range(es.max_row - 4, es.max_row + 1)]
    assert stats[0] == 24 and 0 < stats[1] < 1, stats
    return vals[0], stats


def main(path):
    wb = load_workbook(path)
    assert wb.sheetnames == EXPECTED, wb.sheetnames
    missing = referenced_ids(wb) - source_ids(wb)
    assert not missing, f"unknown source ids: {missing}"
    n = check_strategies(wb["Strategy Playbook"])
    first, stats = check_values(path)
    print(f"OK: {len(EXPECTED)} sheets, {n} strategies, ST01 priority={first:.2f}, stats={stats}")


if __name__ == "__main__":
    main(sys.argv[1])
