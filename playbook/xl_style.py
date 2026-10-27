"""Small openpyxl styling helpers (Arial, header fill, wrap, widths, freeze)."""
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
TITLE = Font(name="Arial", size=14, bold=True)
INPUT = Font(name="Arial", size=10, color="0000FF")
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def title(ws, text, subtitle=None):
    ws["A1"] = text
    ws["A1"].font = TITLE
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = FONT
        ws["A2"].alignment = WRAP
    return 4 if subtitle else 3


def header(ws, row, cols):
    for i, name in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=name)
        c.font, c.fill, c.alignment, c.border = HEAD_FONT, HEAD_FILL, WRAP, BORDER


def write_rows(ws, start, rows):
    for r, values in enumerate(rows, start):
        for i, v in enumerate(values, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.font, c.alignment, c.border = FONT, WRAP, BORDER
    return start + len(rows)


def widths(ws, sizes):
    for i, w in enumerate(sizes, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def finish(ws, header_row, ncols):
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ncols)}{ws.max_row}"
