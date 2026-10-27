"""Assemble the SEO/AEO agentic playbook workbook."""
import sys
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
import xl_style as X
from data_sources import SOURCES
from data_strategies import STRATEGIES
from data_agent import AGENT_EXEC
from data_validation import CHECKS, PENALTIES, ANTIPATTERNS, SCENARIOS
from data_tools import SEO_TOOLS, AEO_TOOLS
from data_misc import TIMELINE, KPIS, PROMPT_SET, EXEC_SUMMARY, README

SUB = "SEO + AEO Agentic Playbook, practice from Aug 2026 to 27 Oct 2026. Evidence grades A-D; see README."


def simple_sheet(wb, name, cols, rows, sizes, subtitle=SUB):
    ws = wb.create_sheet(name)
    r = X.title(ws, name, subtitle)
    X.header(ws, r, cols)
    X.write_rows(ws, r + 1, rows)
    X.widths(ws, sizes)
    X.finish(ws, r, len(cols))
    return ws


def readme_sheet(wb):
    ws = wb.active
    ws.title = "README"
    r = X.title(ws, "SEO + AEO Agentic Playbook (Aug-Oct 2026)", "Prepared 2026-10-27. Legend: blue text = editable estimate inputs; black = formulas or research text.")
    X.header(ws, r, ["Topic", "Notes"])
    X.write_rows(ws, r + 1, README)
    X.widths(ws, [22, 140])


def exec_sheet(wb):
    cols = ["#", "Finding", "Sources"]
    ws = simple_sheet(wb, "Executive Summary", cols, EXEC_SUMMARY, [5, 140, 24])
    r = ws.max_row + 2
    ws.cell(row=r, column=1, value="Portfolio stats (formulas)").font = X.BOLD
    stats = [("Strategies", "=COUNTIF('Strategy Playbook'!A5:A200,\"ST*\")"),
             ("Average win probability", "=AVERAGE('Strategy Playbook'!H5:H200)"),
             ("Strategies with win prob >= 60%", "=COUNTIF('Strategy Playbook'!H5:H200,\">=0.6\")"),
             ("Strategies graded A or B evidence", "=COUNTIF('Strategy Playbook'!G5:G200,\"A\")+COUNTIF('Strategy Playbook'!G5:G200,\"B\")"),
             ("High-risk strategies (risk = 3)", "=COUNTIF('Strategy Playbook'!K5:K200,3)")]
    end = X.write_rows(ws, r + 1, stats)
    ws.cell(row=r + 2, column=2).number_format = "0.0%"
    return end


def strategy_rows():
    rows = []
    for i, s in enumerate(STRATEGIES, 5):
        rows.append([s["id"], s["pillar"], s["name"], s["what"], s["mechanism"], s["evidence"], s["grade"],
                     s["win_prob"], s["impact"], s["effort"], s["risk"],
                     f"=H{i}*I{i}*(4-K{i})/J{i}", s["signal"], s["applies"], s["kpi"], s["sources"]])
    return rows


def strategy_sheet(wb):
    cols = ["ID", "Pillar", "Strategy", "What to do", "Why it works (mechanism)", "Evidence summary", "Evidence grade",
            "Win probability (est.)", "Impact 1-3", "Effort 1-3", "Penalty risk 1-3", "Priority score (formula)",
            "Time to signal", "Applies to", "Primary KPI", "Source IDs"]
    ws = simple_sheet(wb, "Strategy Playbook", cols, strategy_rows(), [6, 11, 30, 60, 55, 45, 9, 12, 8, 8, 10, 12, 12, 28, 32, 18])
    n = 4 + len(STRATEGIES)
    for r in range(5, n + 1):
        ws.cell(row=r, column=8).number_format = "0%"
        ws.cell(row=r, column=12).number_format = "0.00"
        for c in (8, 9, 10, 11):
            ws.cell(row=r, column=c).font = X.INPUT
    ws.conditional_formatting.add(f"L5:L{n}", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"))
    ws.cell(row=n + 2, column=1, value="Blue cells are editable estimates. Priority = win probability x impact x (4 - risk) / effort. Win probability definition in README.").font = X.FONT
    ws.cell(row=n + 2, column=1).fill = X.NOTE_FILL


def agent_sheet(wb):
    names = {s["id"]: s["name"] for s in STRATEGIES}
    cols = ["ID", "Strategy", "Agent role", "Trigger / cadence", "Autonomy level", "Inputs and tools", "Agent workflow steps",
            "Automated validation checks", "Human gate", "Penalty guardrails", "Outputs", "KPIs", "Framework controls (S72)"]
    rows = [[a["id"], names[a["id"]], a["agent"], a["trigger"], a["autonomy"], a["inputs"], a["steps"], a["checks"],
             a["gate"], a["guardrails"], a["outputs"], a["kpis"], a["controls"]] for a in AGENT_EXEC]
    simple_sheet(wb, "Agent Execution", cols, rows, [6, 28, 20, 20, 22, 40, 70, 45, 30, 35, 28, 28, 22])


def validation_sheets(wb):
    simple_sheet(wb, "Validation Checklist", ["ID", "Stage", "Check", "Method", "Tool", "Pass criteria", "Automation", "Frequency", "Enforcement", "Strategies", "Framework controls"],
                 CHECKS, [6, 12, 50, 40, 30, 32, 12, 16, 16, 14, 20])
    simple_sheet(wb, "Penalty Safeguards", ["Policy / risk", "What triggers it", "2026 enforcement evidence", "Agent guardrail", "Detection signal", "Sources"],
                 PENALTIES, [26, 50, 50, 50, 40, 18])
    simple_sheet(wb, "Anti-patterns", ["Anti-pattern", "Evidence against it", "Do instead", "Sources"], ANTIPATTERNS, [36, 60, 50, 16])
    simple_sheet(wb, "Agent Test Scenarios", ["ID", "Situation", "Expected agent behaviour", "Must NOT do", "Verdict rule", "Strategies", "Controls"],
                 SCENARIOS, [6, 50, 55, 35, 30, 12, 14])


def tool_sheets(wb):
    cols = ["Tool", "Category", "Role (track / improve / validate / execute)", "What it does (2026)", "API / MCP", "Pricing (Oct 2026, verify)", "Notes", "Sources"]
    simple_sheet(wb, "SEO Tools", cols, SEO_TOOLS, [34, 22, 18, 60, 30, 30, 36, 14])
    simple_sheet(wb, "AEO Tools", cols, AEO_TOOLS, [34, 22, 18, 60, 24, 36, 36, 14])


def misc_sheets(wb):
    simple_sheet(wb, "Measurement KPIs", ["KPI", "Definition", "Source", "Formula", "Cadence", "Target / benchmark", "Caveat"], KPIS, [26, 50, 30, 24, 12, 40, 30])
    simple_sheet(wb, "Prompt Set Design", ["Element", "Guidance"], PROMPT_SET, [20, 140])
    simple_sheet(wb, "Timeline Aug-Oct 2026", ["Date", "Event", "Implication", "Sources"], TIMELINE, [22, 80, 60, 14])
    simple_sheet(wb, "Sources", ["ID", "Title / finding", "Publisher", "Date", "URL", "Grade"], SOURCES, [6, 90, 30, 18, 60, 7])


def main(out):
    wb = Workbook()
    readme_sheet(wb)
    exec_sheet(wb)
    strategy_sheet(wb)
    agent_sheet(wb)
    validation_sheets(wb)
    tool_sheets(wb)
    misc_sheets(wb)
    wb.save(out)
    print("saved", out)


if __name__ == "__main__":
    main(sys.argv[1])
